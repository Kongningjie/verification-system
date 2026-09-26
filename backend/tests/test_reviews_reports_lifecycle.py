import io
import json
import logging
from pathlib import Path

import pytest
from app.core.config import Settings
from app.core.logging import JsonLogFormatter
from app.db.session import create_engine
from app.services.lifecycle import delete_task
from app.services.storage import task_directories
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from tests.test_checklist_api import create_task
from tests.test_execution_api import confirm


def completed_task(api_client: TestClient, tmp_path: Path) -> tuple[dict, list[dict]]:
    task = create_task(api_client, tmp_path)
    confirm(api_client, task["id"])
    results = api_client.get(
        f"/api/v1/tasks/{task['id']}/results", params={"page_size": 100}
    ).json()["items"]
    return task, results


def test_review_preserves_system_conclusion_and_enforces_comment_rules(
    api_client: TestClient, tmp_path: Path
) -> None:
    task, results = completed_task(api_client, tmp_path)
    critical = next(item for item in results if item["severity"] == "CRITICAL")
    new_value = "FAIL" if critical["final_conclusion"] != "FAIL" else "PASS"
    missing = api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{critical['id']}/review",
        json={"final_conclusion": new_value, "comment": ""},
    )
    assert missing.status_code == 422
    assert missing.json()["code"] == "REVIEW_COMMENT_REQUIRED"

    reviewed = api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{critical['id']}/review",
        json={"final_conclusion": new_value, "comment": "人工核验原始材料后调整。"},
    )
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["reviewer_name"] == "本地用户"
    detail = api_client.get(f"/api/v1/tasks/{task['id']}/results/{critical['id']}").json()
    assert detail["system_conclusion"] == critical["system_conclusion"]
    assert detail["final_conclusion"] == new_value
    assert detail["has_manual_override"] == (new_value != critical["system_conclusion"])
    assert len(detail["review_records"]) == 1


def test_noncritical_fail_to_pass_requires_comment(api_client: TestClient, tmp_path: Path) -> None:
    task, results = completed_task(api_client, tmp_path)
    normal = next(item for item in results if item["severity"] == "NORMAL")
    first = "FAIL" if normal["final_conclusion"] != "FAIL" else "NEEDS_REVIEW"
    response = api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{normal['id']}/review",
        json={"final_conclusion": first, "comment": ""},
    )
    assert response.status_code == 200
    missing = api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{normal['id']}/review",
        json={"final_conclusion": "PASS", "comment": ""},
    )
    assert missing.status_code == 422
    assert missing.json()["code"] == "REVIEW_COMMENT_REQUIRED"


def test_excel_and_json_reports_use_final_conclusion_snapshot(
    api_client: TestClient, tmp_path: Path
) -> None:
    task, results = completed_task(api_client, tmp_path)
    target = next(item for item in results if item["severity"] == "NORMAL")
    changed = "FAIL" if target["final_conclusion"] != "FAIL" else "PASS"
    review = api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{target['id']}/review",
        json={"final_conclusion": changed, "comment": "报告快照测试"},
    )
    assert review.status_code == 200

    json_report = api_client.post(f"/api/v1/tasks/{task['id']}/reports", json={"format": "JSON"})
    assert json_report.status_code == 200, json_report.text
    downloaded = api_client.get(f"/api/v1/tasks/{task['id']}/reports/{json_report.json()['id']}")
    assert downloaded.status_code == 200
    payload = downloaded.json()
    assert set(payload) >= {
        "schema_version",
        "task",
        "files",
        "checklist_snapshot",
        "results",
        "evidence",
        "review_records",
        "execution_metadata",
    }
    exported = next(item for item in payload["results"] if item["id"] == target["id"])
    assert exported["final_conclusion"] == changed
    serialized = json.dumps(payload).casefold()
    assert "api_key" not in serialized
    assert "storage_path" not in serialized
    assert "asset_path" not in serialized
    assert "base64" not in serialized
    assert "prompt_versions" in payload["execution_metadata"]

    second_value = "PASS" if changed != "PASS" else "NEEDS_REVIEW"
    api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{target['id']}/review",
        json={"final_conclusion": second_value, "comment": "报告生成后的再次复核"},
    )
    frozen = api_client.get(f"/api/v1/tasks/{task['id']}/reports/{json_report.json()['id']}").json()
    frozen_result = next(item for item in frozen["results"] if item["id"] == target["id"])
    assert frozen_result["final_conclusion"] == changed

    excel_report = api_client.post(f"/api/v1/tasks/{task['id']}/reports", json={"format": "EXCEL"})
    assert excel_report.status_code == 200, excel_report.text
    workbook_bytes = api_client.get(
        f"/api/v1/tasks/{task['id']}/reports/{excel_report.json()['id']}"
    ).content
    workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True)
    assert workbook.sheetnames == ["Summary", "Results", "Evidence", "ReviewLog", "Metadata"]
    summary_values = {row[0]: row[1] for row in workbook["Summary"].iter_rows(values_only=True)}
    assert summary_values["文件"]
    assert summary_values["规则版本"]
    assert "模型" in summary_values
    rows = list(workbook["Results"].iter_rows(values_only=True))
    headers = list(rows[0])
    final_index = headers.index("最终结论")
    id_index = headers.index("结果 ID")
    result_row = next(row for row in rows[1:] if row[id_index] == target["id"])
    assert result_row[final_index] == second_value
    evidence_rows = list(workbook["Evidence"].iter_rows(values_only=True))
    assert evidence_rows[0] == (
        "证据 ID",
        "核对项 ID",
        "角色",
        "来源类型",
        "来源文件",
        "定位",
        "摘录",
        "哈希",
    )
    assert len(evidence_rows) > 1
    review_rows = list(workbook["ReviewLog"].iter_rows(values_only=True))
    assert len(review_rows) == 3
    assert {row[5] for row in review_rows[1:]} == {"报告快照测试", "报告生成后的再次复核"}


def test_result_pagination_and_manual_filter(api_client: TestClient, tmp_path: Path) -> None:
    task, results = completed_task(api_client, tmp_path)
    target = next(item for item in results if item["severity"] == "NORMAL")
    changed = "FAIL" if target["final_conclusion"] != "FAIL" else "PASS"
    api_client.post(
        f"/api/v1/tasks/{task['id']}/results/{target['id']}/review",
        json={"final_conclusion": changed, "comment": "筛选测试"},
    )
    page = api_client.get(
        f"/api/v1/tasks/{task['id']}/results",
        params={"page": 1, "page_size": 1},
    ).json()
    assert page["total"] == len(results)
    assert len(page["items"]) == 1
    filtered = api_client.get(
        f"/api/v1/tasks/{task['id']}/results",
        params={"has_manual_override": True, "page_size": 100},
    ).json()
    assert filtered["total"] == 1
    assert filtered["items"][0]["id"] == target["id"]


def test_delete_removes_database_and_files(api_client: TestClient, tmp_path: Path) -> None:
    task = create_task(api_client, tmp_path)
    settings = api_client.test_settings  # type: ignore[attr-defined]
    root = task_directories(settings, task["id"])["root"]
    assert root.exists()
    response = api_client.delete(f"/api/v1/tasks/{task['id']}")
    assert response.status_code == 204
    assert not root.exists()
    assert api_client.get(f"/api/v1/tasks/{task['id']}").status_code == 404


def test_delete_completed_task_removes_results_reports_and_disk(
    api_client: TestClient, tmp_path: Path
) -> None:
    task, results = completed_task(api_client, tmp_path)
    generated = api_client.post(
        f"/api/v1/tasks/{task['id']}/reports", json={"format": "JSON"}
    ).json()
    settings = api_client.test_settings  # type: ignore[attr-defined]
    root = task_directories(settings, task["id"])["root"]
    assert root.exists()

    response = api_client.delete(f"/api/v1/tasks/{task['id']}")

    assert response.status_code == 204
    assert not root.exists()
    assert api_client.get(f"/api/v1/tasks/{task['id']}").status_code == 404
    assert (
        api_client.get(f"/api/v1/tasks/{task['id']}/results/{results[0]['id']}").status_code == 404
    )
    assert (
        api_client.get(f"/api/v1/tasks/{task['id']}/reports/{generated['id']}").status_code == 404
    )


def test_file_cleanup_failure_is_explicit_and_retryable(
    api_client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    task = create_task(api_client, tmp_path)
    settings = api_client.test_settings  # type: ignore[attr-defined]
    staged = settings.tasks_root / f".deleting-{task['id']}"
    original = __import__("shutil").rmtree

    def fail(_path):
        raise OSError("locked")

    monkeypatch.setattr("app.services.lifecycle.shutil.rmtree", fail)
    response = api_client.delete(f"/api/v1/tasks/{task['id']}")
    assert response.status_code == 500
    assert response.json()["code"] == "TASK_FILE_CLEANUP_FAILED"
    assert response.json()["details"]["retryable"] is True
    assert staged.exists()
    monkeypatch.setattr("app.services.lifecycle.shutil.rmtree", original)
    retried = api_client.delete(f"/api/v1/tasks/{task['id']}")
    assert retried.status_code == 204
    assert not staged.exists()


def test_download_never_accepts_a_path(api_client: TestClient, tmp_path: Path) -> None:
    task, _ = completed_task(api_client, tmp_path)
    response = api_client.get(f"/api/v1/tasks/{task['id']}/reports/..%2F..%2F.env")
    assert response.status_code in {404, 422}


def test_report_failure_does_not_change_task_status(
    api_client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    task, _ = completed_task(api_client, tmp_path)

    def fail_report(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr("app.services.reports.write_json_report", fail_report)
    response = api_client.post(f"/api/v1/tasks/{task['id']}/reports", json={"format": "JSON"})
    assert response.status_code == 500
    assert set(response.json()) == {"code", "message", "details", "request_id"}
    assert response.json()["code"] == "REPORT_GENERATION_FAILED"
    assert api_client.get(f"/api/v1/tasks/{task['id']}").json()["status"] == "COMPLETED"
    report_id = response.json()["details"]["report_id"]
    unavailable = api_client.get(f"/api/v1/tasks/{task['id']}/reports/{report_id}")
    assert unavailable.status_code == 409
    assert unavailable.json()["code"] == "REPORT_NOT_READY"


def test_database_delete_failure_restores_task_directory(
    api_client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    task = create_task(api_client, tmp_path)
    settings: Settings = api_client.test_settings  # type: ignore[attr-defined]
    root = task_directories(settings, task["id"])["root"]
    engine = create_engine(settings.database_url, connect_args={"check_same_thread": False})
    with Session(engine) as session:
        original_commit = session.commit

        def fail_commit() -> None:
            raise RuntimeError("database unavailable")

        monkeypatch.setattr(session, "commit", fail_commit)
        with pytest.raises(Exception, match="database records"):
            delete_task(session, settings, task["id"])
        monkeypatch.setattr(session, "commit", original_commit)
    engine.dispose()
    assert root.exists()
    assert api_client.get(f"/api/v1/tasks/{task['id']}").status_code == 200


def test_business_log_formatter_always_emits_frozen_fields() -> None:
    record = logging.LogRecord("app.test", logging.INFO, __file__, 1, "完成", (), None)
    payload = json.loads(JsonLogFormatter().format(record))
    assert set(payload) == {
        "timestamp",
        "level",
        "message",
        "request_id",
        "task_id",
        "stage",
        "check_item_id",
        "operation",
        "duration_ms",
        "error_code",
    }
