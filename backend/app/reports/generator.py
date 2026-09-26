import hashlib
import json
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

_SENSITIVE_KEYS = {
    "api_key",
    "openai_api_key",
    "asset_path",
    "storage_path",
    "full_prompt",
    "prompt_text",
    "instructions",
    "image_base64",
}


def safe_cell(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def sanitize_export(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: sanitize_export(item)
            for key, item in value.items()
            if key.casefold() not in _SENSITIVE_KEYS
        }
    if isinstance(value, list):
        return [sanitize_export(item) for item in value]
    return value


def write_json_report(path: Path, snapshot: dict[str, Any]) -> str:
    payload = json.dumps(
        sanitize_export(snapshot), ensure_ascii=False, indent=2, sort_keys=True
    ).encode("utf-8")
    path.write_bytes(payload)
    return hashlib.sha256(payload).hexdigest()


def write_excel_report(path: Path, snapshot: dict[str, Any]) -> str:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    task = snapshot["task"]
    counts = Counter(item["final_conclusion"] for item in snapshot["results"])
    summary_rows = [
        ("任务 ID", task["id"]),
        ("任务名称", task["name"]),
        ("任务状态", task["status"]),
        ("生成时间", snapshot["generated_at"]),
        ("文件数量", len(snapshot["files"])),
        ("文件", ", ".join(item["original_name"] for item in snapshot["files"])),
        ("规则版本", snapshot["execution_metadata"]["ruleset_version"]),
        ("模型", ", ".join(snapshot["execution_metadata"]["models"])),
        *[(f"结论 {key}", value) for key, value in sorted(counts.items())],
    ]
    _append_rows(summary, ["字段", "值"], summary_rows)

    results = workbook.create_sheet("Results")
    _append_rows(
        results,
        [
            "结果 ID",
            "核对项",
            "核对要求",
            "类型",
            "严重等级",
            "系统结论",
            "最终结论",
            "原因码",
            "理由",
            "人工改判",
        ],
        [
            (
                item["id"],
                item["check_name"],
                item["requirement"],
                item["check_type"],
                item["severity"],
                item["system_conclusion"],
                item["final_conclusion"],
                item["reason_code"],
                item["reason"],
                item["has_manual_override"],
            )
            for item in snapshot["results"]
        ],
    )

    evidence = workbook.create_sheet("Evidence")
    _append_rows(
        evidence,
        ["证据 ID", "核对项 ID", "角色", "来源类型", "来源文件", "定位", "摘录", "哈希"],
        [
            (
                item["evidence_id"],
                item["check_item_id"],
                item["role"],
                item["source_category"],
                item.get("source_file_name", ""),
                json.dumps(item["locator"], ensure_ascii=False, sort_keys=True),
                item["excerpt"],
                item["sha256"],
            )
            for item in snapshot["evidence"]
        ],
    )

    reviews = workbook.create_sheet("ReviewLog")
    _append_rows(
        reviews,
        ["复核 ID", "结果 ID", "复核人", "原结论", "新结论", "说明", "时间"],
        [
            (
                item["id"],
                item["result_id"],
                item["reviewer_name"],
                item["previous_conclusion"],
                item["new_conclusion"],
                item["comment"],
                item["created_at"],
            )
            for item in snapshot["review_records"]
        ],
    )

    metadata = workbook.create_sheet("Metadata")
    execution = snapshot["execution_metadata"]
    _append_rows(
        metadata,
        ["字段", "值"],
        [
            ("清单版本", execution["checklist_version"]),
            ("规则版本", execution["ruleset_version"]),
            ("提示词版本", ", ".join(execution["prompt_versions"])),
            ("模型", ", ".join(execution["models"])),
            ("模型运行数", execution["run_count"]),
            ("输入 Token", execution["input_tokens"]),
            ("输出 Token", execution["output_tokens"]),
        ],
    )
    workbook.save(path)
    payload = path.read_bytes()
    return hashlib.sha256(payload).hexdigest()


def _append_rows(sheet, headers: list[str], rows: list[tuple]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width


def generated_at() -> str:
    return datetime.now(UTC).isoformat()
